"""
Excel Export Module - Xuất kết quả ra file Excel định dạng đẹp
"""
import io
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter


# Column definitions — khớp với UI mới
COLUMNS = [
    ("STT", 6),
    ("Tên file", 30),
    ("Số hoá đơn", 18),
    ("Ngày hoá đơn", 14),
    ("Người bán", 35),
    ("Người mua", 35),
    ("Có VAT", 10),
    ("Tiền trước thuế", 18),
    ("Tiền VAT", 18),
    ("Tổng thanh toán", 18),
    ("Trạng thái", 12),
    ("Cảnh báo / Ghi chú", 45),
]


def _get(data: dict, *keys, default=None):
    """Thử lần lượt nhiều tên field, trả về giá trị đầu tiên không None/rỗng."""
    for k in keys:
        v = data.get(k)
        if v is not None and v != "":
            return v
    return default


def export_to_excel(results: List[Dict[str, Any]]) -> io.BytesIO:
    """
    Xuất danh sách kết quả ra file Excel.

    Args:
        results: List of processed invoice data dicts

    Returns:
        BytesIO buffer containing the Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Kết quả bóc tách hoá đơn"

    # --- Styles ---
    header_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Arial", size=10)
    data_alignment = Alignment(vertical="center", wrap_text=True)
    number_alignment = Alignment(horizontal="right", vertical="center")

    warning_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    error_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
    ok_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # --- Title row ---
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    title_cell = ws.cell(row=1, column=1, value="BÁO CÁO BÓC TÁCH HOÁ ĐƠN")
    title_cell.font = Font(name="Arial", bold=True, size=14, color="2F5496")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # --- Header row ---
    header_row = 3
    ws.row_dimensions[header_row].height = 30
    for col_idx, (col_name, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    # --- Data rows ---
    for row_idx, data in enumerate(results, 1):
        excel_row = header_row + row_idx
        ws.row_dimensions[excel_row].height = 22

        # Determine row fill based on status
        status = str(data.get("trang_thai", "") or "")
        row_fill = None
        if status.startswith("CẢNH BÁO"):
            row_fill = warning_fill
        elif status.startswith("LỖI"):
            row_fill = error_fill
        elif status.startswith("OK"):
            row_fill = ok_fill

        # --- Field mapping: hỗ trợ cả upload tab lẫn lark processing ---
        filename      = _get(data, "filename", "_filename", default="")
        so_hoa_don    = _get(data, "so_hoa_don", default="")
        ngay_hoa_don  = _get(data, "ngay_hoa_don", default="")
        # Người bán: thử nhiều tên
        ten_nguoi_ban = _get(data, "ten_nguoi_ban", "ten_nha_cung_cap", default="")
        ten_nguoi_mua = _get(data, "ten_nguoi_mua", default="")

        co_vat_raw = data.get("co_vat")
        if co_vat_raw is True or co_vat_raw == "Có":
            co_vat_str = "Có"
        elif co_vat_raw is False or co_vat_raw == "Không":
            co_vat_str = "Không"
        else:
            co_vat_str = str(co_vat_raw) if co_vat_raw else ""

        # Số tiền: thử tên mới (upload) rồi tên cũ (computed)
        tien_truoc_thue = _get(data, "tien_truoc_thue", "tong_tien_truoc_thue_computed", default=None)
        tien_vat        = _get(data, "tien_vat",        "tien_vat_computed",            default=None)
        tong_thanh_toan = _get(data, "tong_thanh_toan", "tong_thanh_toan_computed",     default=None)

        notes = "; ".join(filter(None, [
            str(data.get("canh_bao") or ""),
            str(data.get("ghi_chu") or ""),
            str(data.get("error") or ""),
        ]))

        row_data = [
            row_idx,          # col 1:  STT
            filename,         # col 2:  Tên file
            so_hoa_don,       # col 3:  Số hoá đơn
            ngay_hoa_don,     # col 4:  Ngày hoá đơn
            ten_nguoi_ban,    # col 5:  Người bán
            ten_nguoi_mua,    # col 6:  Người mua
            co_vat_str,       # col 7:  Có VAT
            tien_truoc_thue,  # col 8:  Tiền trước thuế
            tien_vat,         # col 9:  Tiền VAT
            tong_thanh_toan,  # col 10: Tổng thanh toán
            status,           # col 11: Trạng thái
            notes,            # col 12: Cảnh báo / Ghi chú
        ]

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=excel_row, column=col_idx, value=value if value is not None else "")
            cell.font = data_font
            cell.border = thin_border

            # Number columns (8, 9, 10) - format with thousands separator
            if col_idx in (8, 9, 10) and isinstance(value, (int, float)) and value is not None:
                cell.number_format = '#,##0'
                cell.alignment = number_alignment
            elif col_idx == 1:  # STT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 7:  # Có VAT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = data_alignment

            # Apply status-based fill to status column (col 11)
            if col_idx == 11 and row_fill:
                cell.fill = row_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- Summary row ---
    if results:
        summary_row = header_row + len(results) + 2
        ws.cell(row=summary_row, column=7, value="TỔNG CỘNG:").font = Font(
            name="Arial", bold=True, size=11
        )

        # Sum columns
        for col_idx in (8, 9, 10):
            col_letter = get_column_letter(col_idx)
            start = header_row + 1
            end = header_row + len(results)
            cell = ws.cell(
                row=summary_row, column=col_idx,
                value=f"=SUM({col_letter}{start}:{col_letter}{end})"
            )
            cell.font = Font(name="Arial", bold=True, size=11)
            cell.number_format = '#,##0'
            cell.alignment = number_alignment
            cell.border = Border(top=Side(style="double"), bottom=Side(style="double"))

    # --- Freeze panes ---
    ws.freeze_panes = f"A{header_row + 1}"

    # --- Auto filter ---
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{header_row + len(results)}"

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
